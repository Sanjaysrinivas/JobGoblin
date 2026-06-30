"""Helpers that build tiny in-memory PDF and DOCX files for tests.

Kept out of ``conftest`` so route/service tests can import them directly.
"""

import io

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_pdf_bytes(text: str = "Jane Doe\nSenior Engineer\nPython, FastAPI") -> bytes:
    """Render ``text`` to a one-page PDF using fpdf2 (pure Python)."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in text.split("\n"):
        pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    out = pdf.output()  # bytearray in fpdf2 >= 2.8
    return bytes(out)


def make_docx_bytes(text: str = "Jane Doe\nSenior Engineer\nPython, FastAPI") -> bytes:
    """Build a minimal DOCX with one paragraph per line."""
    from docx import Document

    document = Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
